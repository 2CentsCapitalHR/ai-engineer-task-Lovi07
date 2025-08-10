# red_flag_detector.py
import re
from docx import Document

# -----------------------
# Placeholders for each document type
# -----------------------
PLACEHOLDERS = {
    "adgm-ra-resolution-incorporation-ltd-sole-shareholder-ver-3-0-20170202.docx": [
        r"\[ *[Ii]nsert company name *\]",
        r"\[ *[Ii]nsert date *\]",
        r"\[ *[Ii]nsert name of shareholder *\]",
        r"\[ *[Ii]nsert name of authorised signatories *\]",
        r"\[ *[Ii]nsert name of directors *\]",
    ],
}

# -----------------------
# Mandatory clauses for each document type
# -----------------------
MANDATORY_CLAUSES = {
    "adgm-ra-resolution-incorporation-ltd-sole-shareholder-ver-3-0-20170202.docx": [
        "RESOLVED, that the Company be incorporated in the Abu Dhabi Global Market",
        "hereby appointed and authorised to singly execute all documents and take all necessary and appropriate actions on behalf of the incorporating shareholder to incorporate the Company and is hereby appointed and authorised to execute all documents and take all necessary appropriate actions on behalf of the incorporating shareholder following incorporation.",
        "hereby appointed as director of the Company.",
        "RESOLVED, that the Company duly adopts proposed Articles of Association for the purpose of incorporation of the Company in the Abu Dhabi Global Market.",
    ]
}

# -----------------------
# Helper: Extract text from docx
# -----------------------
def extract_text_from_docx(path):
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

# -----------------------
# Placeholder check
# -----------------------
def check_placeholders(docx_path, doc_type):
    """Check for unfilled placeholders."""
    text = extract_text_from_docx(docx_path)
    patterns = PLACEHOLDERS.get(doc_type, [])
    unfilled = []
    for p in patterns:
        matches = re.findall(p, text)
        if matches:
            unfilled.extend(matches)
    return unfilled

# -----------------------
# Clause check
# -----------------------
def check_missing_clauses(docx_path, doc_type):
    """Check for missing mandatory clauses."""
    text = extract_text_from_docx(docx_path).lower()
    clauses = MANDATORY_CLAUSES.get(doc_type, [])
    missing = []
    for clause in clauses:
        if clause.lower() not in text:
            missing.append(clause)
    return missing

# -----------------------
# LLM-based advanced legal issue detection
# -----------------------
def llm_detect_wrong_jurisdiction(doc_text, doc_name, model, openai_client):
    prompt = f"""
    You are a compliance checker for ADGM legal documents.

    Document: {doc_name}

    Expected Jurisdiction: "Abu Dhabi Global Market" (ADGM Courts)
    Common Wrong Jurisdictions: UAE Federal Courts, Dubai Courts, Abu Dhabi Judicial Department, DIFC Courts.

    Task:
    - Read the document text carefully.
    - Flag any clause or part of the text that mentions a jurisdiction other than ADGM Courts.
    - If no jurisdiction is mentioned, do not flag.
    - Return results strictly in JSON array format:
      [
        {{
          "issue": "Wrong jurisdiction found: '<exact text>'",
          "severity": "High",
          "suggestion": "Replace with 'Abu Dhabi Global Market (ADGM) Courts'."
        }}
      ]
    - If no wrong jurisdiction is found, return an empty list [].

    Document text:
    {doc_text}
    """

    response = openai_client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
    )

    import json
    try:
        return json.loads(response.choices[0].message.content)
    except:
        return []
